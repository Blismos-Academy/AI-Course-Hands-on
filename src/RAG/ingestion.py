from docx import Document
from llama_index.core import Document as LlamaDocument


class DocumentIngestion:

    def load_document(self):

        doc = Document(
            "D:\Agentic AI Project\AI-Course-Hands-on\Data\weather_data.docx"
        )

        text = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)

        document = LlamaDocument(
            text="\n".join(text)
        )

        return [document]